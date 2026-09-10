"""Build the GEEFLIP SQLite database from the source spreadsheets.

The catalog workbook supplies the products, the Keepa export supplies the Amazon
photo for each ASIN, and a previous geeflip.db (if one is lying around) supplies
the winning listings I have already reviewed.

Run directly to rebuild from scratch:
    python geeflip/importer.py --rebuild
"""

from __future__ import annotations

import sqlite3
import sys
from pathlib import Path

PACKAGE_ROOT = Path(__file__).resolve().parent
if str(PACKAGE_ROOT) not in sys.path:
    sys.path.insert(0, str(PACKAGE_ROOT))

from openpyxl import load_workbook

from paths import CATALOG_XLSX, DB_PATH, KEEPA_XLSX
from store import Store, as_float, as_int, clean_listing_title, clean_text, now_iso

CATALOG_COLUMNS = {
    "asin": "ASIN",
    "title": "TITLE",
    "cleaned_title": "CLEANED TITLE",
    "ean": "EAN",
    "upc": "UPC",
    "brand": "Brand",
    "buybox_price": "Buybox (30 days)",
    "sales_rank": "SALES RANK",
    "drops_count": "DROPS (90 DAYS)",
    "restriction": "RESTRICTION",
    "amazon_url": "URL: Amazon",
    "ebay_search_url": "ebay listing",
    "last_enrichment_date": "LAST ENRICHMENT DATE",
}


def _sheet_rows(path: Path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        index = {header: position for position, header in enumerate(headers)}
        for row in rows:
            yield row, index
    finally:
        workbook.close()


def _cell(row: tuple, index: dict[str, int], header: str):
    position = index.get(header)
    if position is None or position >= len(row):
        return None
    return row[position]


def import_catalog(store: Store, path: Path = CATALOG_XLSX) -> int:
    if not path.exists():
        raise FileNotFoundError(f"Catalog workbook not found: {path}")

    batch: list[tuple] = []
    # Row 1 holds the headers, so numbering from 2 keeps these aligned with the
    # row numbers Excel shows and with the scraper's start_row/end_row filters.
    for source_row, (row, index) in enumerate(_sheet_rows(path), start=2):
        batch.append(
            (
                source_row,
                clean_text(_cell(row, index, CATALOG_COLUMNS["asin"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["title"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["cleaned_title"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["ean"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["upc"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["brand"])),
                as_float(_cell(row, index, CATALOG_COLUMNS["buybox_price"])),
                as_int(_cell(row, index, CATALOG_COLUMNS["sales_rank"])),
                as_int(_cell(row, index, CATALOG_COLUMNS["drops_count"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["restriction"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["amazon_url"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["ebay_search_url"])),
                clean_text(_cell(row, index, CATALOG_COLUMNS["last_enrichment_date"])),
            )
        )

    store.replace_products(batch)
    store.set_setting("catalog_imported_at", now_iso())
    store.set_setting("catalog_source", str(path))
    return len(batch)


def import_amazon_images(store: Store, path: Path = KEEPA_XLSX) -> int:
    """Attach the first Keepa image URL to every catalog ASIN it knows."""
    if not path.exists():
        raise FileNotFoundError(f"Keepa workbook not found: {path}")

    matched = 0
    batch: list[tuple[str, str]] = []
    for row, index in _sheet_rows(path):
        asin = clean_text(_cell(row, index, "ASIN")).upper()
        images = clean_text(_cell(row, index, "Image"))
        if not asin or not images:
            continue
        batch.append((images.split(";")[0].strip(), asin))
        if len(batch) >= 500:
            matched += store.apply_amazon_images(batch)
            batch.clear()
    if batch:
        matched += store.apply_amazon_images(batch)

    store.set_setting("images_imported_at", now_iso())
    return matched


LEGACY_WINNER_MAP = {
    "run_id": "run_id",
    "history_key": "history_key",
    "found_at": "found_at",
    "first_won_at": "first_won_at",
    "asin": "asin",
    "title": "title",
    "ean": "ean",
    "brand": "brand",
    "buybox": "buybox",
    "amazon_url": "amazon_url",
    "ebay_url": "ebay_listing_url",
    "ebay_title": "ebay_listing_title",
    "ebay_query": "ebay_listing_query",
    "ebay_image_url": "ebay_image_url",
    "ebay_full_cost": "ebay_full_cost",
    "roi": "roi",
    "seller": "seller",
    "seller_reviews": "seller_reviews",
    "listing_date": "listing_date",
    "search_source": "search_source",
    "seen": "seen",
}


def migrate_legacy(store: Store, legacy_path: Path) -> dict:
    """Carry winners, runs and no-match history over from a pre-V2 database."""
    if not legacy_path.exists():
        return {"winners": 0, "runs": 0, "no_match": 0}

    conn = sqlite3.connect(legacy_path)
    conn.row_factory = sqlite3.Row
    try:
        tables = {
            row["name"]
            for row in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")
        }

        runs: list[dict] = []
        if "scrape_runs" in tables:
            for row in conn.execute("SELECT * FROM scrape_runs ORDER BY id"):
                runs.append(
                    {
                        "id": row["id"],
                        "started_at": row["started_at"],
                        "finished_at": row["finished_at"],
                        "status": row["status"],
                        "filters_json": row["settings_json"],
                        "winners_count": row["winners_count"],
                        "error": row["error"],
                    }
                )
            store.insert_runs(runs)

        winners: list[dict] = []
        if "winning_listings" in tables:
            columns = {
                row[1] for row in conn.execute("PRAGMA table_info(winning_listings)")
            }
            for row in conn.execute("SELECT * FROM winning_listings ORDER BY id"):
                record = {
                    target: row[source]
                    for target, source in LEGACY_WINNER_MAP.items()
                    if source in columns
                }
                record.setdefault("found_at", now_iso())
                record["first_won_at"] = record.get("first_won_at") or record["found_at"]
                for text_column in (
                    "asin", "title", "ean", "brand", "buybox", "amazon_url",
                    "ebay_url", "ebay_title", "ebay_query", "ebay_image_url",
                    "seller", "listing_date", "search_source",
                ):
                    record[text_column] = record.get(text_column) or ""
                record["ebay_title"] = clean_listing_title(record["ebay_title"])
                record["seen"] = int(record.get("seen") or 0)
                # The old schema had a standalone mismatched flag; V2 folds that
                # into the exclusive matched/mismatched verdict.
                if "mismatched" in columns and row["mismatched"]:
                    record["verdict"] = "mismatched"
                    record["reviewed_at"] = record["found_at"]
                    record["seen"] = 1
                winners.append(record)
            store.insert_winner_rows(winners)

        no_match: list[dict] = []
        if "identifier_no_match" in tables:
            no_match = [
                dict(row) for row in conn.execute("SELECT * FROM identifier_no_match")
            ]
            store.save_identifier_no_match(no_match)
    finally:
        conn.close()

    return {"winners": len(winners), "runs": len(runs), "no_match": len(no_match)}


def rebuild(db_path: Path = DB_PATH) -> dict:
    """Drop the database and rebuild it, preserving reviewed winners."""
    legacy_path = db_path.with_name("geeflip.legacy.db")
    if db_path.exists():
        for suffix in ("", "-wal", "-shm"):
            stale = legacy_path.with_name(legacy_path.name + suffix)
            if stale.exists():
                stale.unlink()
        db_path.replace(legacy_path)
        for suffix in ("-wal", "-shm"):
            sidecar = db_path.with_name(db_path.name + suffix)
            if sidecar.exists():
                sidecar.replace(legacy_path.with_name(legacy_path.name + suffix))

    store = Store(db_path)
    try:
        products = import_catalog(store)
        print(f"Catalog: {products:,} products")
        images = import_amazon_images(store)
        print(f"Amazon photos: matched {images:,} ASINs")
        carried = migrate_legacy(store, legacy_path)
        if carried["winners"]:
            print(
                f"Carried over {carried['winners']:,} winners, "
                f"{carried['runs']} runs, {carried['no_match']} no-match records"
            )
        summary = store.stats()
    finally:
        store.close()

    if legacy_path.exists():
        for suffix in ("", "-wal", "-shm"):
            stale = legacy_path.with_name(legacy_path.name + suffix)
            if stale.exists():
                stale.unlink()

    return summary


def ensure_ready(store: Store) -> dict:
    """Populate an empty database on first boot; otherwise leave it alone."""
    if store.product_count() > 0:
        return {"imported": False, **store.stats()}
    products = import_catalog(store)
    images = import_amazon_images(store)
    return {"imported": True, "products": products, "images": images, **store.stats()}


if __name__ == "__main__":
    summary = rebuild()
    print(
        f"GEEFLIP database ready: {summary['products']:,} products, "
        f"{summary['with_image']:,} with Amazon photos, "
        f"{summary['winners']:,} winning listings"
    )
