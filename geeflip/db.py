from __future__ import annotations

import json
import re
import sqlite3
import threading
from datetime import datetime
from pathlib import Path
from urllib.parse import urlsplit, urlunsplit

from openpyxl import load_workbook

GEEFLIP_ROOT = Path(__file__).resolve().parent
PROJECT_ROOT = GEEFLIP_ROOT.parent
DATA_DIR = GEEFLIP_ROOT / "data"
DB_PATH = DATA_DIR / "geeflip.db"
EXCEL_PATH = PROJECT_ROOT / "clean" / "10krows with List 1 6-25.xlsx"
KEEPA_XLSX = PROJECT_ROOT / "clean" / "combined_keepa.xlsx"
WINNERS_HISTORY_JSON = PROJECT_ROOT / "ebay" / "data" / "json" / "winning_listings_history.json"
WINNERS_LIVE_JSON = PROJECT_ROOT / "ebay" / "data" / "json" / "winning_listings_live.json"
NO_MATCH_JSON = PROJECT_ROOT / "ebay" / "data" / "json" / "identifier_no_match_history.json"

DEFAULT_FILTERS = {
    "start_row": None,
    "end_row": None,
    "limit": None,
    "include_brands": [],
    "exclude_brands": [],
    "min_roi_percent": 80.0,
    "max_roi_percent": 300.0,
    "max_listing_age_days": 2,
    "min_seller_reviews": 40,
    "min_sales_rank": None,
    "max_sales_rank": 500_000,
    "min_drop_count": None,
    "min_buybox_price": 40.0,
    "upc_as_well": True,
    "cleaned_title_as_well": False,
    "titles_only": False,
    "image_search": False,
    "skip_previously_won": True,
    "winner_history_retention_days": 3,
    "identifier_no_match_retention_days": 3,
}

SCHEMA = """
CREATE TABLE IF NOT EXISTS products (
    id INTEGER PRIMARY KEY,
    source_row INTEGER NOT NULL UNIQUE,
    title TEXT,
    asin TEXT,
    ean TEXT,
    upc TEXT,
    cleaned_title TEXT,
    sales_rank INTEGER,
    drops_count INTEGER,
    brand TEXT,
    buybox_price REAL,
    amazon_url TEXT,
    ebay_listing TEXT,
    restriction TEXT,
    last_enrichment_date TEXT,
    amazon_image_url TEXT,
    amazon_image_urls TEXT,
    mismatch_count INTEGER NOT NULL DEFAULT 0
);
CREATE INDEX IF NOT EXISTS idx_products_asin ON products(asin);
CREATE INDEX IF NOT EXISTS idx_products_brand ON products(brand);

CREATE TABLE IF NOT EXISTS winning_listings (
    id INTEGER PRIMARY KEY,
    run_id INTEGER,
    found_at TEXT NOT NULL,
    history_key TEXT,
    title TEXT,
    asin TEXT,
    ean TEXT,
    brand TEXT,
    buybox TEXT,
    amazon_url TEXT,
    ebay_full_cost REAL,
    seller TEXT,
    seller_reviews INTEGER,
    listing_date TEXT,
    roi REAL,
    ebay_listing_url TEXT,
    ebay_listing_query TEXT,
    search_source TEXT,
    first_won_at TEXT,
    ebay_listing_title TEXT,
    ebay_image_url TEXT,
    seen INTEGER NOT NULL DEFAULT 0,
    mismatched INTEGER NOT NULL DEFAULT 0
);
CREATE UNIQUE INDEX IF NOT EXISTS idx_winners_history_key
    ON winning_listings(history_key)
    WHERE history_key IS NOT NULL AND history_key != '';
CREATE INDEX IF NOT EXISTS idx_winners_run ON winning_listings(run_id);
CREATE INDEX IF NOT EXISTS idx_winners_found ON winning_listings(found_at);

CREATE TABLE IF NOT EXISTS identifier_no_match (
    query_type TEXT NOT NULL,
    identifier TEXT NOT NULL,
    consecutive_no_matches INTEGER NOT NULL DEFAULT 1,
    no_exact_match_at TEXT NOT NULL,
    PRIMARY KEY (query_type, identifier)
);

CREATE TABLE IF NOT EXISTS scrape_runs (
    id INTEGER PRIMARY KEY,
    started_at TEXT NOT NULL,
    finished_at TEXT,
    status TEXT NOT NULL,
    settings_json TEXT,
    winners_count INTEGER NOT NULL DEFAULT 0,
    error TEXT
);

CREATE TABLE IF NOT EXISTS app_settings (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS meta (
    key TEXT PRIMARY KEY,
    value TEXT NOT NULL
);
"""


def identifier_query_key(identifier_type: str, identifier: object) -> str:
    value = str(identifier or "").strip().casefold()
    return f"{identifier_type.upper()}:{value}" if value else ""


def winner_listing_key(winner: dict) -> str:
    url = str(winner.get("EBAY listing URL") or winner.get("ebay_listing_url") or "").strip()
    if not url:
        return ""
    parts = urlsplit(url)
    item_match = re.search(
        r"/itm/(?:[^/?#]+/)?(\d{9,})",
        parts.path,
        re.IGNORECASE,
    )
    if item_match:
        return f"ebay-item:{item_match.group(1)}"
    normalized_path = parts.path.rstrip("/")
    return urlunsplit(
        (
            parts.scheme.casefold(),
            parts.netloc.casefold(),
            normalized_path,
            "",
            "",
        )
    )


def first_keepa_image_url(value: object) -> str:
    text = identifier_text(value)
    if not text:
        return ""
    return text.split(";")[0].strip()


def identifier_text(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value).rstrip("0").rstrip(".")
    return str(value).strip()


def as_int(value: object) -> int | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    match = re.search(r"-?\d[\d,]*", str(value))
    return int(match.group(0).replace(",", "")) if match else None


def as_float(value: object) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    match = re.search(r"-?\d[\d,]*(?:\.\d+)?", str(value))
    return float(match.group(0).replace(",", "")) if match else None


def cell_date(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return str(value).strip()


class GeeflipStore:
    def __init__(self, path: Path = DB_PATH):
        DATA_DIR.mkdir(parents=True, exist_ok=True)
        self.path = path
        self._lock = threading.RLock()
        self.current_run_id: int | None = None
        self._products_cache: list[dict[str, str]] | None = None
        self._conn = sqlite3.connect(path, check_same_thread=False)
        self._conn.row_factory = sqlite3.Row
        self._conn.execute("PRAGMA journal_mode=WAL")
        self._conn.execute("PRAGMA foreign_keys=ON")
        self._conn.executescript(SCHEMA)
        self._ensure_columns()
        self._conn.commit()

    def _ensure_columns(self) -> None:
        winner_columns = {
            row[1]
            for row in self._conn.execute("PRAGMA table_info(winning_listings)")
        }
        winner_additions = {
            "ebay_listing_title": "ALTER TABLE winning_listings ADD COLUMN ebay_listing_title TEXT",
            "ebay_image_url": "ALTER TABLE winning_listings ADD COLUMN ebay_image_url TEXT",
            "seen": "ALTER TABLE winning_listings ADD COLUMN seen INTEGER NOT NULL DEFAULT 0",
            "mismatched": "ALTER TABLE winning_listings ADD COLUMN mismatched INTEGER NOT NULL DEFAULT 0",
        }
        for name, statement in winner_additions.items():
            if name not in winner_columns:
                self._conn.execute(statement)

        product_columns = {
            row[1]
            for row in self._conn.execute("PRAGMA table_info(products)")
        }
        product_additions = {
            "amazon_image_url": "ALTER TABLE products ADD COLUMN amazon_image_url TEXT",
            "amazon_image_urls": "ALTER TABLE products ADD COLUMN amazon_image_urls TEXT",
            "mismatch_count": "ALTER TABLE products ADD COLUMN mismatch_count INTEGER NOT NULL DEFAULT 0",
        }
        for name, statement in product_additions.items():
            if name not in product_columns:
                self._conn.execute(statement)

    def close(self) -> None:
        with self._lock:
            self._conn.close()

    def get_meta(self, key: str, default: str | None = None) -> str | None:
        with self._lock:
            row = self._conn.execute(
                "SELECT value FROM meta WHERE key = ?",
                (key,),
            ).fetchone()
        return row["value"] if row else default

    def set_meta(self, key: str, value: str) -> None:
        with self._lock:
            self._conn.execute(
                "INSERT INTO meta(key, value) VALUES(?, ?) "
                "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (key, value),
            )
            self._conn.commit()

    def load_filters(self) -> dict:
        with self._lock:
            row = self._conn.execute(
                "SELECT value FROM app_settings WHERE key = 'filters'"
            ).fetchone()
        if not row:
            return dict(DEFAULT_FILTERS)
        try:
            data = json.loads(row["value"])
        except json.JSONDecodeError:
            return dict(DEFAULT_FILTERS)
        merged = dict(DEFAULT_FILTERS)
        merged.update(data)
        return merged

    def save_filters(self, filters: dict) -> dict:
        merged = dict(DEFAULT_FILTERS)
        merged.update(filters)
        with self._lock:
            self._conn.execute(
                "INSERT INTO app_settings(key, value) VALUES('filters', ?) "
                "ON CONFLICT(key) DO UPDATE SET value = excluded.value",
                (json.dumps(merged),),
            )
            self._conn.commit()
        return merged

    def product_count(self) -> int:
        with self._lock:
            row = self._conn.execute("SELECT COUNT(*) AS n FROM products").fetchone()
        return int(row["n"])

    def winner_count(self, *, run_id: int | None = None) -> int:
        with self._lock:
            if run_id is None:
                row = self._conn.execute(
                    "SELECT COUNT(*) AS n FROM winning_listings"
                ).fetchone()
            else:
                row = self._conn.execute(
                    "SELECT COUNT(*) AS n FROM winning_listings WHERE run_id = ?",
                    (run_id,),
                ).fetchone()
        return int(row["n"])

    def catalog_stats(self) -> dict:
        with self._lock:
            row = self._conn.execute(
                """
                SELECT
                    COUNT(*) AS products,
                    COUNT(DISTINCT brand) AS brands,
                    SUM(CASE WHEN ebay_listing IS NOT NULL AND ebay_listing != '' THEN 1 ELSE 0 END) AS with_ebay_url
                FROM products
                """
            ).fetchone()
        return {
            "products": int(row["products"] or 0),
            "brands": int(row["brands"] or 0),
            "with_ebay_url": int(row["with_ebay_url"] or 0),
            "imported_at": self.get_meta("imported_at"),
            "source_file": self.get_meta("source_file"),
        }

    def load_products(self) -> list[dict[str, str]]:
        with self._lock:
            if self._products_cache is not None:
                return self._products_cache
            rows = self._conn.execute(
                """
                SELECT source_row, title, asin, ean, upc, cleaned_title,
                       sales_rank, drops_count, brand, buybox_price,
                       amazon_url, ebay_listing
                FROM products
                ORDER BY source_row
                """
            ).fetchall()
            products: list[dict[str, str]] = []
            for row in rows:
                products.append(
                    {
                        "csv_row": int(row["source_row"]),
                        "title": str(row["title"] or ""),
                        "asin": str(row["asin"] or ""),
                        "ean": str(row["ean"] or ""),
                        "upc": str(row["upc"] or ""),
                        "cleaned_title": str(row["cleaned_title"] or "").strip(),
                        "sales_rank": "" if row["sales_rank"] is None else str(row["sales_rank"]),
                        "drops_count": "" if row["drops_count"] is None else str(row["drops_count"]),
                        "brand": str(row["brand"] or ""),
                        "buybox_price": "" if row["buybox_price"] is None else str(row["buybox_price"]),
                        "amazon_url": str(row["amazon_url"] or ""),
                        "search_url": str(row["ebay_listing"] or "").strip(),
                    }
                )
            self._products_cache = products
            return products

    def delete_identifiers(self, deletion_keys: set[str]) -> int:
        if not deletion_keys:
            return 0
        changed = 0
        with self._lock:
            rows = self._conn.execute("SELECT id, ean, upc FROM products").fetchall()
            for row in rows:
                updates: dict[str, str] = {}
                for column, query_type in (("ean", "EAN"), ("upc", "UPC")):
                    raw = str(row[column] or "")
                    codes = [code.strip() for code in raw.split(",") if code.strip()]
                    remaining = [
                        code
                        for code in codes
                        if identifier_query_key(query_type, code) not in deletion_keys
                    ]
                    if remaining != codes:
                        updates[column] = ", ".join(remaining)
                if not updates:
                    continue
                assignments = ", ".join(f"{column} = ?" for column in updates)
                self._conn.execute(
                    f"UPDATE products SET {assignments} WHERE id = ?",
                    [*updates.values(), row["id"]],
                )
                changed += 1
            self._products_cache = None
            self._conn.commit()
        return changed

    def load_identifier_no_match(self) -> list[dict]:
        with self._lock:
            rows = self._conn.execute(
                """
                SELECT query_type, identifier, consecutive_no_matches, no_exact_match_at
                FROM identifier_no_match
                """
            ).fetchall()
        return [dict(row) for row in rows]

    def save_identifier_no_match(self, records: list[dict]) -> None:
        with self._lock:
            self._conn.execute("DELETE FROM identifier_no_match")
            self._conn.executemany(
                """
                INSERT INTO identifier_no_match(
                    query_type, identifier, consecutive_no_matches, no_exact_match_at
                ) VALUES (?, ?, ?, ?)
                """,
                [
                    (
                        str(record.get("query_type") or ""),
                        str(record.get("identifier") or ""),
                        int(record.get("consecutive_no_matches") or 1),
                        str(record.get("no_exact_match_at") or ""),
                    )
                    for record in records
                    if str(record.get("query_type") or "") and str(record.get("identifier") or "")
                ],
            )
            self._conn.commit()

    def load_winner_history(self) -> list[dict]:
        with self._lock:
            rows = self._conn.execute(
                """
                SELECT * FROM winning_listings
                ORDER BY COALESCE(first_won_at, found_at)
                """
            ).fetchall()
        return [self._row_to_winner(row) for row in rows]

    def save_winner_history(self, records: list[dict]) -> None:
        keys = [
            str(record.get("history_key") or winner_listing_key(record) or "")
            for record in records
        ]
        keys = [key for key in keys if key]
        with self._lock:
            if keys:
                placeholders = ",".join("?" * len(keys))
                self._conn.execute(
                    f"DELETE FROM winning_listings WHERE history_key IS NULL "
                    f"OR history_key NOT IN ({placeholders})",
                    keys,
                )
            else:
                self._conn.execute("DELETE FROM winning_listings")
            for record in records:
                self._upsert_winner_locked(record, run_id=record.get("run_id"))
            self._conn.commit()

    def save_live_winners(self, winners: list[dict]) -> None:
        if self.current_run_id is None:
            return
        with self._lock:
            self._conn.execute(
                "UPDATE scrape_runs SET winners_count = ? WHERE id = ?",
                (len(winners), self.current_run_id),
            )
            self._conn.commit()

    def append_winners(self, winners: list[dict]) -> None:
        with self._lock:
            for winner in winners:
                self._upsert_winner_locked(winner, run_id=self.current_run_id)
            if self.current_run_id is not None:
                row = self._conn.execute(
                    "SELECT COUNT(*) AS n FROM winning_listings WHERE run_id = ?",
                    (self.current_run_id,),
                ).fetchone()
                self._conn.execute(
                    "UPDATE scrape_runs SET winners_count = ? WHERE id = ?",
                    (int(row["n"]), self.current_run_id),
                )
            self._conn.commit()

    def begin_run(self, filters: dict) -> int:
        started = datetime.now().isoformat(timespec="seconds")
        with self._lock:
            cursor = self._conn.execute(
                """
                INSERT INTO scrape_runs(started_at, status, settings_json, winners_count)
                VALUES (?, 'running', ?, 0)
                """,
                (started, json.dumps(filters)),
            )
            self._conn.commit()
            self.current_run_id = int(cursor.lastrowid)
        return self.current_run_id

    def finish_run(
        self,
        status: str,
        *,
        winners_count: int | None = None,
        error: str | None = None,
    ) -> None:
        if self.current_run_id is None:
            return
        finished = datetime.now().isoformat(timespec="seconds")
        with self._lock:
            if winners_count is None:
                row = self._conn.execute(
                    "SELECT COUNT(*) AS n FROM winning_listings WHERE run_id = ?",
                    (self.current_run_id,),
                ).fetchone()
                winners_count = int(row["n"])
            self._conn.execute(
                """
                UPDATE scrape_runs
                SET finished_at = ?, status = ?, winners_count = ?, error = ?
                WHERE id = ?
                """,
                (finished, status, winners_count, error, self.current_run_id),
            )
            self._conn.commit()

    def latest_run(self) -> dict | None:
        with self._lock:
            row = self._conn.execute(
                "SELECT * FROM scrape_runs ORDER BY id DESC LIMIT 1"
            ).fetchone()
        return dict(row) if row else None

    def list_winners(
        self,
        *,
        run_id: int | None = None,
        limit: int = 100,
        offset: int = 0,
    ) -> list[dict]:
        sql = """
            SELECT w.*, r.started_at AS scrape_started_at,
                   p.amazon_image_url AS amazon_image_url,
                   p.mismatch_count AS asin_mismatch_count
            FROM winning_listings w
            LEFT JOIN scrape_runs r ON r.id = w.run_id
            LEFT JOIN products p ON p.id = (
                SELECT id FROM products
                WHERE UPPER(asin) = UPPER(COALESCE(w.asin, ''))
                LIMIT 1
            )
        """
        params: list = []
        if run_id is not None:
            sql += " WHERE w.run_id = ?"
            params.append(run_id)
        sql += " ORDER BY w.found_at DESC, w.id DESC LIMIT ? OFFSET ?"
        params.extend([limit, offset])
        with self._lock:
            rows = self._conn.execute(sql, params).fetchall()
        return [self._row_to_winner(row) for row in rows]

    def set_winner_seen(self, winner_id: int, seen: bool) -> dict | None:
        with self._lock:
            self._conn.execute(
                "UPDATE winning_listings SET seen = ? WHERE id = ?",
                (1 if seen else 0, winner_id),
            )
            self._conn.commit()
            row = self._conn.execute(
                """
                SELECT w.*, r.started_at AS scrape_started_at,
                       p.amazon_image_url AS amazon_image_url,
                       p.mismatch_count AS asin_mismatch_count
                FROM winning_listings w
                LEFT JOIN scrape_runs r ON r.id = w.run_id
                LEFT JOIN products p ON p.id = (
                SELECT id FROM products
                WHERE UPPER(asin) = UPPER(COALESCE(w.asin, ''))
                LIMIT 1
            )
                WHERE w.id = ?
                """,
                (winner_id,),
            ).fetchone()
        return self._row_to_winner(row) if row else None

    def set_winner_mismatched(self, winner_id: int, mismatched: bool) -> dict | None:
        with self._lock:
            current = self._conn.execute(
                "SELECT id, asin, mismatched FROM winning_listings WHERE id = ?",
                (winner_id,),
            ).fetchone()
            if current is None:
                return None
            was = bool(current["mismatched"])
            if was != mismatched:
                self._conn.execute(
                    "UPDATE winning_listings SET mismatched = ? WHERE id = ?",
                    (1 if mismatched else 0, winner_id),
                )
                delta = 1 if mismatched else -1
                self._conn.execute(
                    """
                    UPDATE products
                    SET mismatch_count = MAX(0, COALESCE(mismatch_count, 0) + ?)
                    WHERE UPPER(asin) = UPPER(?)
                    """,
                    (delta, current["asin"] or ""),
                )
                self._conn.commit()
            row = self._conn.execute(
                """
                SELECT w.*, r.started_at AS scrape_started_at,
                       p.amazon_image_url AS amazon_image_url,
                       p.mismatch_count AS asin_mismatch_count
                FROM winning_listings w
                LEFT JOIN scrape_runs r ON r.id = w.run_id
                LEFT JOIN products p ON p.id = (
                SELECT id FROM products
                WHERE UPPER(asin) = UPPER(COALESCE(w.asin, ''))
                LIMIT 1
            )
                WHERE w.id = ?
                """,
                (winner_id,),
            ).fetchone()
        return self._row_to_winner(row) if row else None

    def list_asins(
        self,
        *,
        query: str = "",
        offset: int = 0,
        limit: int = 50,
    ) -> dict:
        query = query.strip()
        where = ""
        params: list = []
        if query:
            like = f"%{query}%"
            where = (
                "WHERE asin LIKE ? OR title LIKE ? OR brand LIKE ? "
                "OR ean LIKE ? OR upc LIKE ?"
            )
            params.extend([like, like, like, like, like])
        with self._lock:
            total = int(
                self._conn.execute(
                    f"SELECT COUNT(*) AS n FROM products {where}",
                    params,
                ).fetchone()["n"]
            )
            rows = self._conn.execute(
                f"""
                SELECT source_row, asin, title, brand, buybox_price, sales_rank,
                       ean, upc, amazon_url, restriction, ebay_listing,
                       amazon_image_url, mismatch_count
                FROM products
                {where}
                ORDER BY source_row
                LIMIT ? OFFSET ?
                """,
                [*params, limit, offset],
            ).fetchall()
        return {
            "total": total,
            "offset": offset,
            "limit": limit,
            "asins": [dict(row) for row in rows],
        }

    def ensure_imported(self, excel_path: Path = EXCEL_PATH) -> dict:
        if self.product_count() > 0:
            return {
                "imported": False,
                "products": self.product_count(),
                "winners": self.winner_count(),
            }
        if not excel_path.exists():
            return {
                "imported": False,
                "skipped": True,
                "products": 0,
                "winners": 0,
            }
        return self.import_excel(excel_path)

    def ensure_keepa_images(self, keepa_path: Path = KEEPA_XLSX) -> dict:
        with self._lock:
            row = self._conn.execute(
                """
                SELECT
                    SUM(CASE WHEN amazon_image_url IS NOT NULL AND amazon_image_url != '' THEN 1 ELSE 0 END) AS with_image,
                    COUNT(*) AS products
                FROM products
                """
            ).fetchone()
        with_image = int(row["with_image"] or 0)
        products = int(row["products"] or 0)
        if self.get_meta("keepa_images_imported") and with_image > 0:
            return {
                "imported": False,
                "with_image": with_image,
                "products": products,
            }
        if not keepa_path.exists():
            return {
                "imported": False,
                "skipped": True,
                "with_image": with_image,
                "products": products,
            }
        return self.import_keepa_images(keepa_path)

    def import_keepa_images(self, keepa_path: Path = KEEPA_XLSX) -> dict:
        if not keepa_path.exists():
            raise FileNotFoundError(f"Keepa workbook not found: {keepa_path}")

        workbook = load_workbook(keepa_path, read_only=True, data_only=True)
        updated = 0
        seen = 0
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, ())]
            index = {header: i for i, header in enumerate(headers)}
            asin_idx = index.get("ASIN")
            image_idx = index.get("Image")
            if asin_idx is None or image_idx is None:
                raise ValueError(f"{keepa_path.name} needs ASIN and Image columns")

            batch: list[tuple[str, str, str]] = []
            with self._lock:
                for row in rows:
                    asin = identifier_text(
                        row[asin_idx] if asin_idx < len(row) else ""
                    ).upper()
                    images = identifier_text(
                        row[image_idx] if image_idx < len(row) else ""
                    )
                    if not asin or not images:
                        continue
                    seen += 1
                    batch.append((first_keepa_image_url(images), images, asin))
                    if len(batch) >= 500:
                        updated += self._apply_keepa_image_batch(batch)
                        batch.clear()
                        print(f"Keepa images: matched {updated:,} catalog ASINs")
                if batch:
                    updated += self._apply_keepa_image_batch(batch)
                self._products_cache = None
                self.set_meta(
                    "keepa_images_imported",
                    datetime.now().isoformat(timespec="seconds"),
                )
                self._conn.commit()
        finally:
            workbook.close()

        with self._lock:
            with_image = int(
                self._conn.execute(
                    """
                    SELECT COUNT(*) AS n FROM products
                    WHERE amazon_image_url IS NOT NULL AND amazon_image_url != ''
                    """
                ).fetchone()["n"]
            )
        print(f"Keepa images ready: {with_image:,} products have Amazon photos")
        return {
            "imported": True,
            "keepa_rows": seen,
            "updated": updated,
            "with_image": with_image,
        }

    def _apply_keepa_image_batch(self, batch: list[tuple[str, str, str]]) -> int:
        before = self._conn.total_changes
        self._conn.executemany(
            """
            UPDATE products
            SET amazon_image_url = ?, amazon_image_urls = ?
            WHERE UPPER(asin) = ?
            """,
            batch,
        )
        return max(self._conn.total_changes - before, 0)

    def import_excel(self, excel_path: Path = EXCEL_PATH) -> dict:
        if not excel_path.exists():
            raise FileNotFoundError(f"Catalog workbook not found: {excel_path}")

        workbook = load_workbook(excel_path, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(value or "").strip() for value in next(rows, ())]
            index = {header: i for i, header in enumerate(headers)}
            required = ("TITLE", "ASIN")
            missing = [name for name in required if name not in index]
            if missing:
                raise ValueError(
                    f"{excel_path.name} is missing columns: {', '.join(missing)}"
                )

            batch: list[tuple] = []
            count = 0
            with self._lock:
                self._products_cache = None
                self._conn.execute("DELETE FROM products")
                for source_row, row in enumerate(rows, start=2):
                    batch.append(
                        (
                            source_row,
                            identifier_text(self._col(row, index, "TITLE")),
                            identifier_text(self._col(row, index, "ASIN")),
                            identifier_text(self._col(row, index, "EAN")),
                            identifier_text(self._col(row, index, "UPC")),
                            identifier_text(self._col(row, index, "CLEANED TITLE")),
                            as_int(self._col(row, index, "SALES RANK")),
                            as_int(self._col(row, index, "DROPS (90 DAYS)")),
                            identifier_text(self._col(row, index, "Brand")),
                            as_float(self._col(row, index, "Buybox (30 days)")),
                            identifier_text(self._col(row, index, "URL: Amazon")),
                            identifier_text(self._col(row, index, "ebay listing")),
                            identifier_text(self._col(row, index, "RESTRICTION")),
                            cell_date(self._col(row, index, "LAST ENRICHMENT DATE")),
                        )
                    )
                    if len(batch) >= 1000:
                        self._insert_product_batch(batch)
                        count += len(batch)
                        batch.clear()
                        print(f"Imported {count:,} catalog rows into SQLite")
                if batch:
                    self._insert_product_batch(batch)
                    count += len(batch)
                self.set_meta("imported_at", datetime.now().isoformat(timespec="seconds"))
                self.set_meta("source_file", str(excel_path))
                self._conn.commit()
        finally:
            workbook.close()

        winners_imported = 0
        if self.winner_count() == 0:
            winners_imported = self._import_winner_json()
        no_match_imported = 0
        if not self.load_identifier_no_match():
            no_match_imported = self._import_no_match_json()

        print(f"Catalog ready: {count:,} products in {self.path}")
        return {
            "imported": True,
            "products": count,
            "winners": winners_imported,
            "no_match": no_match_imported,
        }

    def _insert_product_batch(self, batch: list[tuple]) -> None:
        self._conn.executemany(
            """
            INSERT INTO products(
                source_row, title, asin, ean, upc, cleaned_title,
                sales_rank, drops_count, brand, buybox_price, amazon_url,
                ebay_listing, restriction, last_enrichment_date
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            batch,
        )

    @staticmethod
    def _col(row: tuple, index: dict[str, int], name: str):
        position = index.get(name)
        if position is None or position >= len(row):
            return None
        return row[position]

    def _upsert_winner_locked(self, winner: dict, *, run_id: object = None) -> None:
        history_key = str(winner.get("history_key") or winner_listing_key(winner) or "") or None
        found_at = str(
            winner.get("found_at")
            or winner.get("first_won_at")
            or datetime.now().isoformat(timespec="seconds")
        )
        first_won_at = str(winner.get("first_won_at") or found_at)
        stored_run_id = winner.get("run_id") if winner.get("run_id") is not None else run_id
        ebay_listing_title = (
            winner.get("EBAY listing title")
            or winner.get("ebay_listing_title")
            or ""
        )
        ebay_image_url = (
            winner.get("EBAY image URL")
            or winner.get("ebay_image_url")
            or ""
        )
        values = (
            stored_run_id,
            found_at,
            history_key,
            winner.get("title") or "",
            winner.get("ASIN") or winner.get("asin") or "",
            winner.get("EAN") or winner.get("ean") or "",
            winner.get("Brand") or winner.get("brand") or "",
            "" if winner.get("BUYBOX") is None and winner.get("buybox") is None
            else str(winner.get("BUYBOX") if winner.get("BUYBOX") is not None else winner.get("buybox")),
            winner.get("AMAZON URL") or winner.get("amazon_url") or "",
            as_float(winner.get("EBAY full cost") if "EBAY full cost" in winner else winner.get("ebay_full_cost")),
            winner.get("SELLER") or winner.get("seller") or "",
            as_int(winner.get("SELLER REVIEWS") if "SELLER REVIEWS" in winner else winner.get("seller_reviews")),
            winner.get("LISTING DATE") or winner.get("listing_date") or "",
            as_float(winner.get("ROI") if "ROI" in winner else winner.get("roi")),
            winner.get("EBAY listing URL") or winner.get("ebay_listing_url") or "",
            winner.get("ebay listing query") or winner.get("ebay_listing_query") or "",
            winner.get("search_source") or "",
            first_won_at,
            ebay_listing_title,
            ebay_image_url,
        )
        existing = None
        if history_key:
            existing = self._conn.execute(
                "SELECT id FROM winning_listings WHERE history_key = ?",
                (history_key,),
            ).fetchone()
        if existing:
            self._conn.execute(
                """
                UPDATE winning_listings SET
                    run_id = COALESCE(?, run_id),
                    title = ?,
                    asin = ?,
                    ean = ?,
                    brand = ?,
                    buybox = ?,
                    amazon_url = ?,
                    ebay_full_cost = ?,
                    seller = ?,
                    seller_reviews = ?,
                    listing_date = ?,
                    roi = ?,
                    ebay_listing_url = ?,
                    ebay_listing_query = ?,
                    search_source = ?,
                    ebay_listing_title = CASE WHEN ? != '' THEN ? ELSE ebay_listing_title END,
                    ebay_image_url = CASE WHEN ? != '' THEN ? ELSE ebay_image_url END
                WHERE id = ?
                """,
                (
                    stored_run_id,
                    values[3],
                    values[4],
                    values[5],
                    values[6],
                    values[7],
                    values[8],
                    values[9],
                    values[10],
                    values[11],
                    values[12],
                    values[13],
                    values[14],
                    values[15],
                    values[16],
                    ebay_listing_title,
                    ebay_listing_title,
                    ebay_image_url,
                    ebay_image_url,
                    existing["id"],
                ),
            )
            return
        self._conn.execute(
            """
            INSERT INTO winning_listings(
                run_id, found_at, history_key, title, asin, ean, brand, buybox,
                amazon_url, ebay_full_cost, seller, seller_reviews, listing_date,
                roi, ebay_listing_url, ebay_listing_query, search_source, first_won_at,
                ebay_listing_title, ebay_image_url, seen
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
            """,
            values,
        )

    @staticmethod
    def _row_to_winner(row: sqlite3.Row) -> dict:
        keys = set(row.keys())
        return {
            "id": row["id"],
            "run_id": row["run_id"],
            "found_at": row["found_at"],
            "first_won_at": row["first_won_at"],
            "scrape_started_at": row["scrape_started_at"] if "scrape_started_at" in keys else None,
            "history_key": row["history_key"],
            "title": row["title"] or "",
            "ASIN": row["asin"] or "",
            "EAN": row["ean"] or "",
            "Brand": row["brand"] or "",
            "BUYBOX": row["buybox"] or "",
            "AMAZON URL": row["amazon_url"] or "",
            "EBAY full cost": row["ebay_full_cost"],
            "SELLER": row["seller"] or "",
            "SELLER REVIEWS": row["seller_reviews"],
            "LISTING DATE": row["listing_date"] or "",
            "ROI": row["roi"],
            "EBAY listing URL": row["ebay_listing_url"] or "",
            "EBAY listing title": (row["ebay_listing_title"] if "ebay_listing_title" in keys else "") or "",
            "EBAY image URL": (row["ebay_image_url"] if "ebay_image_url" in keys else "") or "",
            "AMAZON image URL": (row["amazon_image_url"] if "amazon_image_url" in keys else "") or "",
            "seen": bool(row["seen"]) if "seen" in keys else False,
            "mismatched": bool(row["mismatched"]) if "mismatched" in keys else False,
            "asin_mismatch_count": (
                int(row["asin_mismatch_count"] or 0)
                if "asin_mismatch_count" in keys
                else 0
            ),
            "ebay listing query": row["ebay_listing_query"] or "",
            "search_source": row["search_source"] or "",
        }

    def _import_winner_json(self) -> int:
        records: list[dict] = []
        for path in (WINNERS_HISTORY_JSON, WINNERS_LIVE_JSON):
            if not path.exists():
                continue
            try:
                payload = json.loads(path.read_text(encoding="utf-8"))
            except (OSError, json.JSONDecodeError):
                continue
            if isinstance(payload, list):
                records.extend(item for item in payload if isinstance(item, dict))
        if not records:
            return 0
        with self._lock:
            for record in records:
                self._upsert_winner_locked(record)
            self._conn.commit()
        return len(records)

    def _import_no_match_json(self) -> int:
        if not NO_MATCH_JSON.exists():
            return 0
        try:
            payload = json.loads(NO_MATCH_JSON.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            return 0
        if not isinstance(payload, list):
            return 0
        records = [item for item in payload if isinstance(item, dict)]
        self.save_identifier_no_match(records)
        return len(records)
