"""Update buy box and sales rank on existing workbook rows from a Keepa export."""

from __future__ import annotations

import os
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from lib.paths import COMBINED_XLSX, PROJECT_ROOT, create_workbook_backup

KEEPA_XLSX = PROJECT_ROOT / "clean" / "combined_keepa.xlsx"
BUYBOX_HEADER = "Buybox (30 days)"
SALES_RANK_HEADER = "SALES RANK"


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def normalize_header(value: object) -> str:
    return " ".join(cell_text(value).casefold().split())


def has_value(value: object) -> bool:
    return value is not None and str(value).strip() != ""


def find_column(headers: list[str], candidates: tuple[str, ...]) -> int | None:
    expected = {normalize_header(candidate) for candidate in candidates}
    for index, header in enumerate(headers):
        if normalize_header(header) in expected:
            return index
    return None


def load_keepa_updates(path: Path) -> dict[str, dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows, ())]
        asin_idx = find_column(headers, ("ASIN",))
        buybox_idx = find_column(
            headers,
            ("Buy Box: Current", "Buy Box: 90 days avg.", BUYBOX_HEADER),
        )
        rank_idx = find_column(
            headers,
            ("Sales Rank: Current", SALES_RANK_HEADER),
        )
        buybox_avg_idx = find_column(headers, ("Buy Box: 90 days avg.",))
        if asin_idx is None:
            raise ValueError(f"{path.name} is missing an ASIN column")

        updates: dict[str, dict] = {}
        for row in rows:
            asin = cell_text(row[asin_idx] if asin_idx < len(row) else "").upper()
            if not asin:
                continue
            buybox = row[buybox_idx] if buybox_idx is not None and buybox_idx < len(row) else None
            if not has_value(buybox) and buybox_avg_idx is not None and buybox_avg_idx < len(row):
                buybox = row[buybox_avg_idx]
            rank = row[rank_idx] if rank_idx is not None and rank_idx < len(row) else None
            payload = {}
            if has_value(buybox):
                payload["buybox"] = buybox
            if has_value(rank):
                payload["sales_rank"] = rank
            if payload:
                updates[asin] = payload
        return updates
    finally:
        workbook.close()


def update_workbook(path: Path, updates: dict[str, dict]) -> tuple[int, int]:
    workbook = load_workbook(path)
    sheet = workbook["main"] if "main" in workbook.sheetnames else workbook.active
    headers = [str(cell.value or "").strip() for cell in sheet[1]]
    asin_idx = find_column(headers, ("ASIN",))
    buybox_idx = find_column(headers, (BUYBOX_HEADER, "BUYBOX"))
    rank_idx = find_column(headers, (SALES_RANK_HEADER, "Sales Rank: Current"))
    if asin_idx is None:
        workbook.close()
        raise ValueError(f"{path.name} is missing an ASIN column")
    if buybox_idx is None or rank_idx is None:
        workbook.close()
        missing = []
        if buybox_idx is None:
            missing.append(BUYBOX_HEADER)
        if rank_idx is None:
            missing.append(SALES_RANK_HEADER)
        raise ValueError(f"{path.name} is missing columns: {', '.join(missing)}")

    updated_rows = 0
    matched_asins: set[str] = set()
    for row_number in range(2, sheet.max_row + 1):
        asin = cell_text(sheet.cell(row=row_number, column=asin_idx + 1).value).upper()
        payload = updates.get(asin)
        if not payload:
            continue
        changed = False
        if "buybox" in payload:
            sheet.cell(row=row_number, column=buybox_idx + 1, value=payload["buybox"])
            changed = True
        if "sales_rank" in payload:
            sheet.cell(row=row_number, column=rank_idx + 1, value=payload["sales_rank"])
            changed = True
        if changed:
            updated_rows += 1
            matched_asins.add(asin)

    temporary_path = path.with_name(f".{path.stem}.tmp{path.suffix}")
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, path)
    finally:
        workbook.close()
        temporary_path.unlink(missing_ok=True)
    return updated_rows, len(matched_asins)


def main() -> int:
    keepa_path = KEEPA_XLSX
    output_path = COMBINED_XLSX
    if not keepa_path.exists():
        print(f"Keepa file not found: {keepa_path}")
        return 1
    if not output_path.exists():
        print(f"Workbook not found: {output_path}")
        return 1

    updates = load_keepa_updates(keepa_path)
    backup_path = create_workbook_backup(output_path)
    print(f"Backed up current workbook to {backup_path.name}")
    updated_rows, matched_asins = update_workbook(output_path, updates)
    print(f"Kept all existing rows in {output_path.name}")
    print(f"Updated buy box and sales rank on {updated_rows} rows ({matched_asins} ASINs)")
    print(f"Keepa ASINs with no matching workbook row: {len(updates) - matched_asins}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
